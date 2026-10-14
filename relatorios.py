from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from atendimentos import atendimentos_do_mes

def gerar_relatorio_mensal(ano, mes, arquivo_saida=None):
    dados = atendimentos_do_mes(ano, mes)
    arquivo_saida = arquivo_saida or f"relatorio_{ano}_{mes:02d}.xlsx"
    
    wbook = Workbook()
    ws = wbook.active
    ws.title = f"{mes:02d}-{ano}"
    
    cabecalhos = ["Nome", "CPF", "Celular", "Data", "Valor Pago (R$)"]
    ws.append(cabecalhos)
    for col, _ in enumerate(cabecalhos, 1):
        c = ws.cell(row=1, column=col)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2f5597")
        c.alignment = Alignment(horizontal="center")
        
    total = 0.0
    for nome, cpf, celular, data, valor in dados:
       try:
           d,m, a = data.split("-")
           data = f"{d}/{m}/{a}"
       except ValueError:
           pass
       ws.append([nome, cpf, celular, data, valor])
       total += valor
        
    ws.append([])
    linha_total = ws.max_row
    c1 = ws.cell(row=linha_total, column=4, value="TOTAL:")
    c1.font = Font(bold=True)
    c2 = ws.cell(row=linha_total, column=5, value=total)
    c2.font = Font(bold=True)
    for col in "ABCDE":
        ws.column_dimensions[col].width = 20
        
    wbook.save(arquivo_saida)
    return arquivo_saida, total